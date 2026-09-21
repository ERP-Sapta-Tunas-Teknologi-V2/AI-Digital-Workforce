import pymupdf
import pymupdf4llm
import re
import tempfile
import subprocess
import platform
from pathlib import Path
from pptx import Presentation
from ingestion.image import describe_image

def _get_libreoffice_command():
    if platform.system() == "Windows":
        return r"C:\Program Files\LibreOffice\program\soffice.exe"
    return "libreoffice"

def _remove_conversion_artifacts(markdown):
    replacements = {
        "Â": "",
        "â€™": "'",
        "â€œ": '"',
        "â€": '"',
        "â€“": "-",
        "â€”": "—",
        "â€¦": "…",
    }

    for old, new in replacements.items():
        markdown = markdown.replace(old, new)

    markdown = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]", "", markdown)
    markdown = re.sub(r"[ \t]+", " ", markdown)
    markdown = re.sub(r"\n{3,}", "\n\n", markdown)

    return markdown

def extract_images(page):
    images = []

    for image in page.get_images(full=True):
        xref = image[0]
        data = page.parent.extract_image(xref)
        images.append(data["image"])

    return images

def docx_to_pdf(docx_path, pdf_path):
    print("Converting .docx to .pdf")
    output_dir = pdf_path.parent

    with tempfile.TemporaryDirectory() as profile_dir:
        profile_uri = Path(profile_dir).as_uri()

        result = subprocess.run(
            [
                _get_libreoffice_command(),
                "--headless",
                f"-env:UserInstallation={profile_uri}",
                "--convert-to", "pdf",
                "--outdir", str(output_dir),
                str(docx_path)
            ],
            capture_output=True,
            text=True,
            timeout=120
        )

    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice conversion failed: {result.stderr}")

    generated_pdf = output_dir / f"{docx_path.stem}.pdf"

    if not generated_pdf.exists():
        raise RuntimeError(f"Expected output not found: {generated_pdf}")

    if generated_pdf != pdf_path:
        generated_pdf.rename(pdf_path)

def pdf_to_md(pdf_path):
    print("Converting .pdf to .md")
    pages = []

    with pymupdf.open(pdf_path) as pdf:
        for page_number, page in enumerate(pdf):
            print(f"\nPage {page_number + 1}")
            markdown = pymupdf4llm.to_markdown(pdf, pages=[page_number])

            # # image-to-text
            # for image_bytes in extract_images(page):
            #     print("image-to-text")
            #     markdown += "\n\n" + describe_image(image_bytes)

            pages.append({"page": page_number + 1, "markdown": markdown})
            
    return pages

def pptx_to_md(pptx_path):
    prs = Presentation(pptx_path)
    pages = []

    for page_number, slide in enumerate(prs.slides, 1):
        texts = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())

        if texts:
            texts[0] = f"## {texts[0]}"

        markdown = "\n\n".join(texts)
        pages.append({"page": page_number, "markdown": markdown})

    return pages

def clean_md(markdown):
    markdown = _remove_conversion_artifacts(markdown)
    markdown = markdown.replace("**", "")  # Remove bold
    markdown = re.sub(r"^# (?!#)", "## ", markdown, flags=re.MULTILINE)  # h1 (#) to h2 (##)
    markdown = re.sub(re.compile("<.*?>"), " ", markdown)  # Remove html tags
    markdown = markdown.replace("`", "")  # Remove `
    return markdown

def preprocessing(path):
    if path.suffix.lower() == ".docx":
        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / f"{path.stem}.pdf"
            docx_to_pdf(path, pdf_path)
            pages = pdf_to_md(pdf_path)

    elif path.suffix.lower() == ".pdf":
        pages = pdf_to_md(path)

    elif path.suffix.lower() == ".pptx":
        pages = pptx_to_md(path)
        
    else:
        raise ValueError(f"Unsupported file type: {path.suffix}")

    print("Cleaning .md")
    for page in pages:
        page["markdown"] = clean_md(page["markdown"])

    return pages