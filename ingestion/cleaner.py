import pymupdf
import pymupdf4llm
import re
import tempfile
import subprocess
import platform
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from ingestion.image import describe_image

def _get_libreoffice_command():
    if platform.system() == "Windows":
        return r"C:\Program Files\LibreOffice\program\soffice.exe"
    return "libreoffice"

def _remove_conversion_artifacts(markdown):
    replacements = {
        "Â": "",
        "â€™": "'",
        "â€œ": '"',
        "â€": '"',
        "â€“": "-",
        "â€”": "—",
        "â€¦": "…",
    }

    for old, new in replacements.items():
        markdown = markdown.replace(old, new)

    markdown = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F]", "", markdown)
    markdown = re.sub(r"[ \t]+", " ", markdown)
    markdown = re.sub(r"\n{3,}", "\n\n", markdown)

    return markdown

def extract_images(page):
    images = []

    for image in page.get_images(full=True):
        xref = image[0]
        data = page.parent.extract_image(xref)
        images.append(data["image"])

    return images

def docx_to_pdf(docx_path, pdf_path):
    print("Converting .docx to .pdf")
    output_dir = pdf_path.parent

    with tempfile.TemporaryDirectory() as profile_dir:
        profile_uri = Path(profile_dir).as_uri()

        result = subprocess.run(
            [
                _get_libreoffice_command(),
                "--headless",
                f"-env:UserInstallation={profile_uri}",
                "--convert-to", "pdf",
                "--outdir", str(output_dir),
                str(docx_path)
            ],
            capture_output=True,
            text=True,
            timeout=120
        )

    if result.returncode != 0:
        raise RuntimeError(f"LibreOffice conversion failed: {result.stderr}")

    generated_pdf = output_dir / f"{docx_path.stem}.pdf"

    if not generated_pdf.exists():
        raise RuntimeError(f"Expected output not found: {generated_pdf}")

    if generated_pdf != pdf_path:
        generated_pdf.rename(pdf_path)

def pdf_to_md(pdf_path):
    print("Converting .pdf to .md")
    pages = []

    with pymupdf.open(pdf_path) as pdf:
        for page_number, page in enumerate(pdf):
            print(f"\nPage {page_number + 1}")
            markdown = pymupdf4llm.to_markdown(pdf, pages=[page_number])

            # # image-to-text
            # for image_bytes in extract_images(page):
            #     print("image-to-text")
            #     markdown += "\n\n" + describe_image(image_bytes)

            pages.append({"page": page_number + 1, "markdown": markdown})
            
    return pages

def _table_to_markdown(table):
    rows = list(table.rows)
    if not rows:
        return ""

    lines = []
    header_cells = [" ".join(cell.text.split()).replace("|", "\\|") for cell in rows[0].cells]
    lines.append("| " + " | ".join(header_cells) + " |")
    lines.append("| " + " | ".join(["---"] * len(header_cells)) + " |")

    for row in rows[1:]:
        row_cells = [" ".join(cell.text.split()).replace("|", "\\|") for cell in row.cells]
        lines.append("| " + " | ".join(row_cells) + " |")

    return "\n".join(lines)

def _shape_texts(shape):
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            yield from _shape_texts(s)
    elif getattr(shape, "has_table", False) and shape.has_table:
        table_md = _table_to_markdown(shape.table)
        if table_md.strip():
            yield table_md
    elif getattr(shape, "has_text_frame", False) and shape.text_frame and shape.text_frame.text.strip():
        yield shape.text_frame.text.strip()

def pptx_to_md(pptx_path):
    prs = Presentation(pptx_path)
    pages = []

    for page_number, slide in enumerate(prs.slides, 1):
        texts = []

        for shape in slide.shapes:
            for text in _shape_texts(shape):
                if text and text.strip():
                    texts.append(text.strip())

        if getattr(slide, "has_notes_slide", False) and slide.has_notes_slide:
            notes_tf = getattr(slide.notes_slide, "notes_text_frame", None)
            if notes_tf and notes_tf.text.strip():
                texts.append(notes_tf.text.strip())

        if texts:
            if not texts[0].startswith("|"):
                texts[0] = f"## {texts[0]}"

        markdown = "\n\n".join(texts)
        pages.append({"page": page_number, "markdown": markdown})

    return pages

def clean_md(markdown):
    markdown = _remove_conversion_artifacts(markdown)
    markdown = markdown.replace("**", "")  # Remove bold
    markdown = re.sub(r"^# (?!#)", "## ", markdown, flags=re.MULTILINE)  # h1 (#) to h2 (##)
    markdown = markdown.replace("`", "")  # Remove `

    # Remove html tags
    markdown = re.sub(r"</[^>]+>\s+(?=[a-zA-Z])", lambda m: m.group(0).replace(" ", ""), markdown)
    markdown = re.sub(r"\s+<[^>]+>", lambda m: m.group(0).lstrip(), markdown)
    markdown = re.sub(r"<[^>]+>", "", markdown).strip()
    
    return markdown

def preprocessing(path):
    if path.suffix.lower() == ".docx":
        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / f"{path.stem}.pdf"
            docx_to_pdf(path, pdf_path)
            pages = pdf_to_md(pdf_path)

    elif path.suffix.lower() == ".pdf":
        pages = pdf_to_md(path)

    elif path.suffix.lower() == ".pptx":
        pages = pptx_to_md(path)
        
    else:
        raise ValueError(f"Unsupported file type: {path.suffix}")

    print("Cleaning .md")
    for page in pages:
        page["markdown"] = clean_md(page["markdown"])

    return pages